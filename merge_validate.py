from __future__ import annotations

import csv
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook

FIELDS = ["id", "name", "email", "amount", "status"]
ISSUE_FIELDS = ["source", "row_number", "id", "reason"]
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def _header(value) -> str:
    return "_".join(str(value or "").strip().lower().split())


def read_rows(paths: Iterable[Path]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for path in sorted((Path(p) for p in paths), key=lambda p: p.name.lower()):
        suffix = path.suffix.lower()
        if suffix == ".csv":
            with path.open("r", encoding="utf-8-sig", newline="") as f:
                reader = csv.reader(f)
                values = list(reader)
        elif suffix == ".xlsx":
            wb = load_workbook(path, read_only=True, data_only=True)
            values = list(wb.active.iter_rows(values_only=True))
        else:
            continue
        if not values:
            continue
        headers = [_header(v) for v in values[0]]
        for row_number, values_row in enumerate(values[1:], start=2):
            row = {headers[i]: "" if i >= len(values_row) or values_row[i] is None else str(values_row[i]) for i in range(len(headers))}
            row["__source"] = path.name
            row["__row_number"] = row_number
            rows.append(row)
    return rows


def normalize_row(row: dict[str, object], source: str | None = None, row_number: int | None = None) -> dict[str, object]:
    amount_raw = str(row.get("amount", "")).strip().replace(",", "")
    amount = ""
    amount_invalid = False
    try:
        if not amount_raw:
            raise InvalidOperation
        amount = f"{Decimal(amount_raw):.2f}"
    except (InvalidOperation, ValueError):
        amount_invalid = True
    status_raw = str(row.get("status", "")).strip().lower()
    status_map = {"new": "New", "pending": "Pending", "paid": "Paid"}
    return {
        "id": str(row.get("id", "")).strip(),
        "name": " ".join(str(row.get("name", "")).split()).title(),
        "email": str(row.get("email", "")).strip().lower(),
        "amount": amount,
        "status": status_map.get(status_raw, status_raw.title()),
        "__source": source if source is not None else str(row.get("__source", "")),
        "__row_number": int(row_number if row_number is not None else row.get("__row_number", 0) or 0),
        "__amount_invalid": amount_invalid,
    }


def validate_row(row: dict[str, object]) -> list[str]:
    issues: list[str] = []
    if not row.get("id"):
        issues.append("missing id")
    if not row.get("name"):
        issues.append("missing name")
    email = str(row.get("email", ""))
    if not email:
        issues.append("missing email")
    elif not EMAIL_RE.match(email):
        issues.append("invalid email")
    if bool(row.get("__amount_invalid")):
        issues.append("invalid amount")
    return issues


def deduplicate(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    ordered = sorted(rows, key=lambda r: (str(r.get("__source", "")).lower(), int(r.get("__row_number", 0) or 0), str(r.get("id", ""))))
    winners: dict[tuple, dict[str, object]] = {}
    for row in ordered:
        row_id = str(row.get("id", ""))
        key = ("id", row_id) if row_id else ("row", str(row.get("__source", "")), int(row.get("__row_number", 0) or 0))
        winners.setdefault(key, row)
    return sorted(winners.values(), key=lambda r: (str(r.get("id", "")), str(r.get("email", "")), str(r.get("__source", ""))))


def write_clean_csv(rows: list[dict[str, object]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in FIELDS})


def write_clean_xlsx(rows: list[dict[str, object]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Clean Data"
    ws.append(FIELDS)
    for row in rows:
        ws.append([row.get(field, "") for field in FIELDS])
    wb.save(path)


def write_issues_report(rows: list[dict[str, object]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=ISSUE_FIELDS)
        writer.writeheader()
        for row in rows:
            for reason in validate_row(row):
                writer.writerow({"source": row.get("__source", ""), "row_number": row.get("__row_number", ""), "id": row.get("id", ""), "reason": reason})
