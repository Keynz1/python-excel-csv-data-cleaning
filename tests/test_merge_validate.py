import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook

try:
    import merge_validate as _merge
except ModuleNotFoundError:
    _merge = None


def impl():
    assert _merge is not None, "merge_validate module is not implemented yet"
    return _merge


def make_csv(path: Path, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(rows)


def make_xlsx(path: Path, rows):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(path)


def test_read_rows_accepts_mixed_csv_xlsx_and_normalizes_headers(tmp_path):
    csv_path = tmp_path / "a.csv"
    xlsx_path = tmp_path / "b.xlsx"
    make_csv(csv_path, [[" ID ", "NAME", "Email", "Amount", "Status"], ["1", "alice smith", "A@EXAMPLE.COM", "12.5", "paid"]])
    make_xlsx(xlsx_path, [["id", " name ", "EMAIL", "amount", "status"], ["2", "Bob Lee", "bob@example.com", "7", "NEW"]])
    rows = impl().read_rows([xlsx_path, csv_path])
    assert len(rows) == 2
    assert rows[0]["id"] == "1"
    assert rows[0]["__source"] == "a.csv"
    assert rows[1]["id"] == "2"


def test_normalize_and_validate_flags_bad_values():
    row = impl().normalize_row({"id": " 3 ", "name": " jane doe ", "email": "not-an-email", "amount": "oops", "status": "pending"}, "sample.csv", 4)
    assert row["id"] == "3"
    assert row["name"] == "Jane Doe"
    assert row["email"] == "not-an-email"
    assert row["amount"] == ""
    assert row["status"] == "Pending"
    issues = impl().validate_row(row)
    assert "invalid email" in issues
    assert "invalid amount" in issues


def test_validate_flags_missing_required_fields():
    row = impl().normalize_row({"id": "", "name": "", "email": "", "amount": "10", "status": "new"}, "sample.csv", 5)
    issues = impl().validate_row(row)
    assert "missing id" in issues
    assert "missing name" in issues
    assert "missing email" in issues


def test_deduplicate_is_input_order_independent_and_uses_stable_source_winner():
    a = impl().normalize_row({"id": "1", "name": "Alice A", "email": "a@example.com", "amount": "10", "status": "new"}, "a.csv", 2)
    b = impl().normalize_row({"id": "1", "name": "Alice B", "email": "b@example.com", "amount": "20", "status": "paid"}, "b.xlsx", 2)
    first = impl().deduplicate([b, a])
    second = impl().deduplicate([a, b])
    assert first == second
    assert first[0]["name"] == "Alice A"


def read_csv_rows(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def read_xlsx_rows(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    values = list(wb.active.iter_rows(values_only=True))
    headers = [str(v) for v in values[0]]
    return [{headers[i]: "" if v is None else str(v) for i, v in enumerate(row)} for row in values[1:]]


def test_clean_csv_xlsx_match_and_issues_have_source_row(tmp_path):
    rows = [
        impl().normalize_row({"id": "1", "name": "Alice", "email": "alice@example.com", "amount": "10", "status": "paid"}, "a.csv", 2),
        impl().normalize_row({"id": "2", "name": "Bob", "email": "bad", "amount": "oops", "status": "new"}, "b.xlsx", 3),
    ]
    csv_path = tmp_path / "clean.csv"
    xlsx_path = tmp_path / "clean.xlsx"
    issues_path = tmp_path / "issues.csv"
    mod = impl()
    mod.write_clean_csv(rows, csv_path)
    mod.write_clean_xlsx(rows, xlsx_path)
    mod.write_issues_report(rows, issues_path)
    assert read_csv_rows(csv_path) == read_xlsx_rows(xlsx_path)
    issues = read_csv_rows(issues_path)
    assert {i["reason"] for i in issues} == {"invalid email", "invalid amount"}
    assert all(i["source"] == "b.xlsx" and i["row_number"] == "3" for i in issues)
